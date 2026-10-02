from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt
from fpdf import FPDF

from backend.utils.text import sanitize_text

ASSETS_DIR = Path(__file__).resolve().parents[2] / "assets"
LOGO_PATH = ASSETS_DIR / "logo.png"


def _split_lines(text: str):
    return [line.strip() for line in sanitize_text(text).split("\n")]


def _is_heading(line: str) -> bool:
    return bool(re.match(r"^(\d+\.|[A-Z][A-Z\s&-]{4,})$", line.strip()))


def format_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")


def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(7)

    if LOGO_PATH.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO_PATH), width=Inches(1.8))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for line in _split_lines(text):
        if not line:
            continue
        if line == doc_type.upper():
            continue
        if _is_heading(line):
            p = doc.add_paragraph()
            r = p.add_run(line)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
        else:
            p = doc.add_paragraph(line)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


    supplied_terms = [item.strip() for item in terms.split(";") if item.strip()]
    if supplied_terms:
        doc.add_paragraph()
        heading = doc.add_paragraph()
        run = heading.add_run("SUPPLIED TERMS SUMMARY")
        run.bold = True
        run.font.name = "Times New Roman"
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "#"
        table.rows[0].cells[1].text = "Term"
        for index, term in enumerate(supplied_terms, start=1):
            cells = table.add_row().cells
            cells[0].text = str(index)
            cells[1].text = term

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase | AI-generated draft | Review before signing").font.size = Pt(8)

    out = BytesIO()
    doc.save(out)
    return out.getvalue()


class LegalEasePDF(FPDF):
    def __init__(self, doc_type: str):
        super().__init__()
        self.doc_type = doc_type
        self.set_auto_page_break(auto=True, margin=18)

    def header(self):
        if LOGO_PATH.exists():
            try:
                self.image(str(LOGO_PATH), x=85, y=7, w=40)
                self.set_y(23)
            except Exception:
                self.set_y(10)
        else:
            self.set_y(10)
        self.set_font("Helvetica", "B", 11)
        self.cell(0, 8, self.doc_type.upper(), align="C")
        self.ln(9)

    def footer(self):
        self.set_y(-13)
        self.set_font("Helvetica", "I", 7)
        self.cell(0, 5, "LegalEase | AI-generated draft | Review before signing", align="C")


def format_pdf(text: str, doc_type: str) -> bytes:
    pdf = LegalEasePDF(doc_type)
    pdf.add_page()
    pdf.set_margins(16, 32, 16)
    pdf.set_font("Helvetica", size=10)

    for line in _split_lines(text):
        if not line or line == doc_type.upper():
            continue
        if _is_heading(line):
            pdf.set_font("Helvetica", "B", 11)
            pdf.multi_cell(0, 6, line)
            pdf.ln(1)
            pdf.set_font("Helvetica", size=10)
        else:
            # Built-in PDF fonts are intentionally used so the project works without shipping proprietary fonts.
            ascii_line = line.encode("latin-1", "replace").decode("latin-1")
            pdf.multi_cell(0, 5, ascii_line)
            pdf.ln(1)

    return bytes(pdf.output())
